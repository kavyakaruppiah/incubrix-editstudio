import io
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import BinaryIO, List, Union


class DocumentExtractionError(Exception):
    pass


class DocumentService:
    """Extract clean text content from PDF, Word (.docx), and plain text documents."""

    def extract_text(self, file_content: Union[bytes, BinaryIO], filename: str) -> str:
        """
        Extract plain text from uploaded file buffer based on extension.
        Supports: .pdf, .docx, .doc, .txt, .md
        """
        if isinstance(file_content, bytes):
            stream = io.BytesIO(file_content)
        else:
            stream = io.BytesIO(file_content.read())

        ext = Path(filename).suffix.lower()

        if ext == ".pdf":
            return self._extract_pdf(stream)
        elif ext == ".docx":
            return self._extract_docx(stream)
        elif ext in (".txt", ".md", ".srt", ".vtt"):
            return self._extract_plain_text(stream)
        elif ext == ".doc":
            # Binary doc format fallback: try plain text extraction of printable characters
            return self._extract_fallback_binary(stream)
        else:
            raise DocumentExtractionError(
                f"Unsupported document format '{ext}'. Please upload a PDF (.pdf), Word document (.docx), or text file (.txt)."
            )

    def _extract_pdf(self, stream: io.BytesIO) -> str:
        try:
            from pypdf import PdfReader
            reader = PdfReader(stream)
            text_parts = []
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text_parts.append(extracted)
            full_text = "\n\n".join(text_parts).strip()
            if not full_text:
                raise DocumentExtractionError("PDF file contains no extractable text (it may be a scanned image).")
            return self._clean_text(full_text)
        except DocumentExtractionError:
            raise
        except Exception as e:
            raise DocumentExtractionError(f"Failed to parse PDF document: {str(e)}")

    def _extract_docx(self, stream: io.BytesIO) -> str:
        # First try python-docx
        try:
            import docx
            doc = docx.Document(stream)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text.strip())
            full_text = "\n\n".join(paragraphs).strip()
            if full_text:
                return self._clean_text(full_text)
        except Exception:
            pass

        # Fallback pure-python XML extraction from docx zip
        try:
            stream.seek(0)
            with zipfile.ZipFile(stream) as zf:
                xml_content = zf.read("word/document.xml")
            tree = ET.fromstring(xml_content)
            # Find all text elements <w:t>
            texts = []
            for elem in tree.iter():
                if elem.tag.endswith("}t") and elem.text:
                    texts.append(elem.text)
            full_text = " ".join(texts).strip()
            if not full_text:
                raise DocumentExtractionError("Word document contains no extractable text.")
            return self._clean_text(full_text)
        except Exception as e:
            raise DocumentExtractionError(f"Failed to parse Word document (.docx): {str(e)}")

    def _extract_plain_text(self, stream: io.BytesIO) -> str:
        stream.seek(0)
        raw_bytes = stream.read()
        for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                decoded = raw_bytes.decode(encoding).strip()
                return self._clean_text(decoded)
            except UnicodeDecodeError:
                continue
        raise DocumentExtractionError("Unable to decode text file. Ensure it is encoded in UTF-8.")

    def _extract_fallback_binary(self, stream: io.BytesIO) -> str:
        stream.seek(0)
        raw = stream.read()
        # Extract ASCII strings of length >= 4
        printable = re.findall(rb"[\x20-\x7E\r\n]{4,}", raw)
        decoded = [p.decode("latin-1", errors="ignore") for p in printable]
        text = "\n".join(decoded).strip()
        if not text:
            raise DocumentExtractionError("Could not extract readable text from legacy .doc file. Please save as .docx or .pdf.")
        return self._clean_text(text)

    def _clean_text(self, text: str) -> str:
        """Clean extra whitespace and normalize quotes."""
        # Replace non-breaking spaces
        text = text.replace("\xa0", " ")
        # Replace smart quotes with standard
        text = text.replace("“", '"').replace("”", '"').replace("‘", "'").replace("’", "'")
        # Collapse multiple blank lines
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()


document_service = DocumentService()
